"""What leaves the building. The buyer gets documents only; the mapping key stays internal.

While the redaction log (original → replacement) exists, the output is pseudonymized, not anonymized,
so it is never packed next to the documents (edge case 4.1).
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path

from .registry import Result, redaction_log


FORMATS = {"pdf": ("PDF", "application/pdf"),
           "docx": ("Word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
           "txt": ("Text", "text/plain")}
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _paragraphs(text: str) -> list[list[str]]:
    """Blank lines separate paragraphs; single line breaks inside one are kept."""
    return [p.split("\n") for p in re.split(r"\n[ \t]*\n", _CONTROL.sub("", text).strip()) if p.strip()]


# A Unicode serif font if the machine has one (packages.txt / the Dockerfile install DejaVu); otherwise the
# built-in Times, which only knows Windows-1252, so the rest is spelled out ("→" becomes "->").
_FONT_PATHS = ("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", "/usr/share/fonts/dejavu/DejaVuSerif.ttf",
               "/usr/share/fonts/TTF/DejaVuSerif.ttf", "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
               "/Library/Fonts/Times New Roman.ttf")
_SPELL = str.maketrans({"→": "->", "←": "<-", "↔": "<->", "≤": "<=", "≥": ">=", "≠": "!=", "≈": "~", "✓": "[x]",
                        "✔": "[x]", "✗": "[ ]", "−": "-", "‐": "-", "‑": "-", "\u2009": " ", "\u202f": " ",
                        "\u200b": "", "\ufeff": ""})


def _pdf_font() -> tuple[str, callable]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "LLBody" in pdfmetrics.getRegisteredFontNames():
        return "LLBody", lambda t: t.translate(_SPELL)
    for path in _FONT_PATHS:
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont("LLBody", path))
                return "LLBody", lambda t: t.translate(_SPELL)
            except Exception:  # noqa: BLE001 — a broken font file just means the fallback
                continue
    return "Times-Roman", lambda t: t.translate(_SPELL).encode("cp1252", "replace").decode("cp1252")


def to_pdf(docs: list[tuple[str, str]]) -> bytes:
    """Each document starts on a new page. No author or source metadata is written."""
    from xml.sax.saxutils import escape

    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=inch, rightMargin=inch, topMargin=inch,
                            bottomMargin=inch, title="Anonymized document", author="", subject="", creator="",
                            producer="", keywords="")
    font, clean = _pdf_font()
    body = ParagraphStyle("body", fontName=font, fontSize=11, leading=15, spaceAfter=9)
    head = ParagraphStyle("head", fontName=font, fontSize=8, leading=10, spaceAfter=16,
                          textColor=HexColor("#6b6b6b"))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(HexColor("#6b6b6b"))
        canvas.drawRightString(LETTER[0] - inch, 0.6 * inch, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    story = []
    for i, (name, text) in enumerate(docs):
        if i:
            story.append(PageBreak())
        story.append(Paragraph(escape(clean(name)), head))
        for lines in _paragraphs(text):
            story.append(Paragraph("<br/>".join(escape(clean(l)) for l in lines), body))
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def to_docx(docs: list[tuple[str, str]]) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    props = d.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category"):
        setattr(props, attr, "")
    props.title = "Anonymized document"
    for i, (name, text) in enumerate(docs):
        if i:
            d.add_page_break()
        run = d.add_paragraph().add_run(name)
        run.font.size, run.font.color.rgb = Pt(8), RGBColor(0x6B, 0x6B, 0x6B)
        for lines in _paragraphs(text):
            p = d.add_paragraph()
            for j, line in enumerate(lines):
                r = p.add_run(line)
                if j < len(lines) - 1:
                    r.add_break()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_file(docs: list[tuple[str, str]], fmt: str = "txt") -> bytes:
    if fmt == "pdf":
        return to_pdf(docs)
    if fmt == "docx":
        return to_docx(docs)
    return "\n\n".join(f"{'=' * 72}\n{name}\n{'=' * 72}\n\n{out}" for name, out in docs).encode("utf-8") \
        if len(docs) > 1 else docs[0][1].encode("utf-8")


def buyer_zip(names: list[str], outputs: list[str], fmt: str = "txt") -> bytes:
    """Anonymized documents only — no redaction log, no mapping key."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, out in zip(names, outputs):
            zf.writestr(f"{Path(name).stem}.anonymized.{fmt}", render_file([(name, out)], fmt))
    return buf.getvalue()


def internal_log_csv(res: Result, only: list[str] | None = None) -> str:
    """The mapping key (original → replacement). Internal and access-controlled; never shipped to a buyer."""
    rows = redaction_log(res)
    if only:
        rows = [r for r in rows if any(n in r["documents"] for n in only)]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()) if rows else ["original"])
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()
