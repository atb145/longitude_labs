"""Regression tests for the Finance pack + general rules (policies/_base.yaml).

Every edge case in "Anonymization Edge Cases — Longitude Labs" (Oct 14, 2026) has a test, plus real-world
documents: the bundled sample case and Microsoft's 8-K of 2022-01-18 (Activision Blizzard).
Run:  python -m pytest -q
"""
import io
import pathlib
import re

import pytest

from engine import Session, consistency_check, redaction_log
from engine.extract import extract_document
from engine.registry import round_figure

ROOT = pathlib.Path(__file__).resolve().parent.parent
FIX = ROOT / "tests" / "fixtures"


def run(docs, **settings):
    s = Session.from_texts(docs)
    s.settings.update(settings)
    return s, s.run()


def entity(res, text):
    return next((e for e in res.entities if any(text.lower() in m[3].lower() for m in e.mentions)), None)


def flags(res, rule):
    return [f for f in res.flags if f.rule == rule]


@pytest.fixture(scope="module")
def edge():
    return run([("edge_cases.txt", (FIX / "edge_cases.txt").read_text())])


@pytest.fixture(scope="module")
def msft():
    data = (FIX / "msft_8k_2022-01-18.htm").read_bytes()
    s = Session.from_files([("msft_8k.htm", data)])
    return s, s.run()


# ---------------------------------------------------------------- 1. Names & public entities
def test_1_1_public_company_that_is_a_party_is_redacted(edge):
    # "We advised Apex Energy…" — Apex is the client, so the name goes even though it's listed.
    _, res = edge
    apex = entity(res, "Apex Energy")
    assert apex is not None and apex.status == "auto" and apex.rule == "names.party_by_context"
    assert "Apex" not in res.outputs[0]


def test_1_1_public_company_cited_for_facts_goes_to_party_check():
    docs = [("d.txt", "Market context: Apex Energy Inc. (NYSE: APXE) reports 35M customers in its 10-K.")]
    _, res = run(docs)
    apex = entity(res, "Apex Energy")
    assert apex.status == "pending" and apex.rule == "names.public_company_party_check"


def test_1_1_public_company_mode_keep_and_redact():
    docs = [("d.txt", "Market context: Apex Energy Inc. (NYSE: APXE) reports 35M customers in its 10-K.")]
    _, kept = run(docs, public_companies="keep")
    assert entity(kept, "Apex Energy").status == "public"
    _, red = run(docs, public_companies="redact")
    assert "Apex" not in red.outputs[0]


def test_known_parties_are_always_redacted():
    s = Session.from_texts([("d.txt", "Contoso met Fabrikam twice. CONTOSO's board approved the Fabrikam deal.")])
    s.set_parties([{"name": "Contoso Holdings", "type": "ORGANIZATION"}, {"name": "Fabrikam", "type": "ORGANIZATION"}])
    out = s.run().outputs[0]
    assert "Contoso" not in out and "CONTOSO" not in out and "Fabrikam" not in out


def test_domain_linking_catches_company_name():
    s, res = run([("d.txt", "Quote from Northwind leadership. Contact ir@northwindtraders.com or northwind.com.")])
    assert "Northwind" not in res.outputs[0]


def test_1_2_eponyms_are_kept(edge):
    _, res = edge
    out = res.outputs[0]
    for term in ("Kaplan-Meier curve", "Black-Scholes value", "Hodgkin lymphoma"):
        assert term in out, term


def test_1_3_citations_and_statutes_are_kept(edge):
    _, res = edge
    out = res.outputs[0]
    assert "SEC v. Texas Gulf Sulphur Co., 401 F.2d 833 (2d Cir. 1968)" in out
    assert "15 U.S.C. § 78j(b)" in out and "Rule 10b-5" in out


def test_1_own_matter_without_reporter_citation_is_still_redacted():
    _, res = run([("01_complaint.txt", (ROOT / "samples" / "01_complaint.txt").read_text())])
    assert "Mercer" not in res.outputs[0]


# ---------------------------------------------------------------- 2. Quasi-identifiers
def test_2_1_rare_profile_is_flagged(edge):
    _, res = edge
    assert any("34-year-old" in f.text for f in flags(res, "quasi.combination"))


def test_2_2_role_place_event_is_flagged(edge):
    _, res = edge
    assert any("Greenwich quant fund" in f.text for f in flags(res, "quasi.combination"))


def test_2_3_shifted_dates_keep_intervals():
    txt = "Bought on March 15, 2022 and March 21, 2022; flagged on March 23, 2022."
    _, res = run([("d.txt", txt)], date_mode="shifted")
    out = res.outputs[0]
    assert "March 15, 2022" not in out and "2022" not in out or "March" not in out
    from datetime import datetime
    got = [datetime.strptime(d, "%B %d, %Y") for d in re.findall(r"[A-Z][a-z]+ \d{1,2}, \d{4}", out)]
    assert [(b - a).days for a, b in zip(got, got[1:])] == [6, 2]


def test_2_3_relative_and_month_year_modes():
    txt = "Bought on March 15, 2022; flagged on March 23, 2022."
    _, rel = run([("d.txt", txt)], date_mode="relative")
    assert "Day 0" in rel.outputs[0] and "Day 8" in rel.outputs[0]
    _, my = run([("d.txt", txt)])
    assert my.outputs[0].count("March 2022") == 2


def test_2_4_headline_figures_are_rounded(edge):
    _, res = edge
    out = res.outputs[0]
    assert "~40,000 shares" in out and "~$5M" in out
    assert "41,000" not in out and "$4.75M" not in out


@pytest.mark.parametrize("raw,expected", [("$4.75M", "~$5M"), ("$68.7 billion", "~$70 billion"),
                                          ("$750,000", "~$800,000"), ("41,000 shares", "~40,000 shares")])
def test_2_4_round_figure(raw, expected):
    assert round_figure(raw) == expected


def test_alt_group_inference_is_flagged(edge):
    _, res = edge
    assert flags(res, "flag.group_inference")


# ---------------------------------------------------------------- 3. Files & images
def test_3_1_docx_hidden_text_is_scanned():
    import docx
    d = docx.Document()
    d.add_paragraph("The memo concerns the Northwind Pharma launch.")
    d.sections[0].header.paragraphs[0].text = "Prepared for Dana Whitfield"
    d.core_properties.author = "jsmith@firm.com"
    try:
        d.add_comment(d.paragraphs[0].runs, text="Per Dana (Northwind BD): don't share with FDA team",
                      author="Dana Whitfield")
    except AttributeError:
        pass  # older python-docx without comments
    buf = io.BytesIO()
    d.save(buf)
    ex = extract_document("memo.docx", buf.getvalue())
    assert "[Header] Prepared for Dana Whitfield" in ex.text
    assert "jsmith@firm.com" in ex.text
    s = Session.from_files([("memo.docx", buf.getvalue())])
    out = s.run().outputs[0]
    assert "jsmith@firm.com" not in out and "Whitfield" not in out


def test_3_1_textless_pdf_page_blocks_release():
    from engine.registry import release_check
    s, res = run([("a.txt", "Nothing to see.")])
    blockers = release_check(res, [[{"level": "block", "message": "Page 2 has no text layer"}]], True)
    assert any("no text layer" in b for b in blockers)


# ---------------------------------------------------------------- 4. Release & rights
def test_4_1_log_is_not_in_buyer_download():
    import zipfile
    from engine.export import buyer_zip
    s, res = run([("01_complaint.txt", (ROOT / "samples" / "01_complaint.txt").read_text())])
    z = zipfile.ZipFile(io.BytesIO(buyer_zip(res.names, res.outputs)))
    assert not any("log" in n.lower() for n in z.namelist())


def test_4_2_privileged_content_is_flagged(edge):
    _, res = edge
    assert any("Privileged" in f.text for f in flags(res, "rights.privileged"))


def test_4_3_nickname_is_linked_for_review(edge):
    _, res = edge
    bob = entity(res, "Bob")
    vance = entity(res, "Robert Vance")
    assert bob is not None and bob.status == "pending" and bob.proposed == vance.label


def test_4_3_quoted_nickname_is_linked_automatically():
    _, res = run([("d.txt", 'Jonathan "Jack" Mercer runs the fund. Later, Jack moved the money.')])
    assert "Jack" not in res.outputs[0]


def test_release_blocked_until_rights_and_review():
    s, res = run([("d.txt", (FIX / "edge_cases.txt").read_text())])
    assert s.release_blockers(res)
    s.rights_confirmed = True
    for k in res.queue:
        s.decide(k, "accept")
    assert s.release_blockers(s.run()) == []


# ---------------------------------------------------------------- real documents
def test_msft_8k(msft):
    s, res = msft
    out = res.outputs[0]
    for leak in ("Microsoft", "MICROSOFT", "Activision", "Dolliver", "91-1144442", "001-37845", "882-8080",
                 "microsoft.com", "One Microsoft Way", "$68.7 billion", "MSFT"):
        assert leak not in out or leak in "".join(f.text for f in res.flags if f.status == "pending"), leak
    for keep in ("Securities Exchange Act of 1934", "Rule 425", "17 CFR 230.425", "Exhibit 99.1", "Regulation FD"):
        assert keep in out, keep
    assert not any(e.canonical.lower() in ("press release", "state", "irs") for e in res.entities)
    cc = consistency_check(res)
    assert cc["all_consistent"]


def test_microsoft_press_release_is_caught():
    txt = ("PRESS RELEASE\n\nMicrosoft Announces Partnership with Contoso Ltd.\n\nREDMOND, Wash. — Microsoft Corp. "
           "(Nasdaq: MSFT) today announced a partnership. Microsoft's revenue grew. "
           "Visit news.microsoft.com or email press@microsoft.com.")
    _, res = run([("pr.txt", txt)], public_companies="redact")
    out = res.outputs[0]
    assert "Microsoft" not in out.replace("[", "")
    assert not any("announces" in e.canonical.lower() or e.canonical.lower() == "press release" for e in res.entities)


def test_sample_case_consistent_and_leak_free():
    docs = [(p.name, p.read_text()) for p in sorted((ROOT / "samples").glob("*.txt"))]
    _, res = run(docs)
    cc = consistency_check(res)
    assert cc["all_consistent"] and cc["leak_free"]
    assert len({e.label for e in res.entities if e.type == "PERSON" and ":nick:" not in e.key}) >= 3


def test_delete_entity_removes_false_positive():
    s, res = run([("d.txt", "The BLCT orders were routed through Meridian Clearing yesterday.")])
    m = entity(res, "Meridian")
    s.delete(m.key)
    res2 = s.run()
    assert entity(res2, "Meridian") is None and "Meridian Clearing" in res2.outputs[0]
    assert any(r["decision"] == "deleted" for r in redaction_log(res2))


# ---------------------------------------------------------------- labels
def test_role_labels_from_context():
    memo = (FIX / "advisory_memo.txt").read_text()
    _, res = run([("memo.txt", memo)], role_labels=True)
    labels = {e.canonical: e.label for e in res.entities}
    assert labels["Northwind Traders Inc."] == "Client"
    assert labels["Tailspin Logistics"] == "Target"
    assert labels["Fabrikam Holdings"] == "Counterparty"
    assert "Northwind" not in res.outputs[0] and "Client" in res.outputs[0]


def test_role_labels_caption_two_defendants_and_no_gaps():
    docs = [(p.name, p.read_text()) for p in sorted((ROOT / "samples").glob("*.txt"))]
    _, res = run(docs, role_labels=True)
    labels = {e.canonical: e.label for e in res.entities}
    assert labels["Jonathan R. Mercer"] == "Defendant A" and labels["Harborview Capital Partners LLC"] == "Defendant B"
    assert labels["Daniel K. Okafor"] == "Respondent"
    companies = sorted(l for l in labels.values() if l.startswith("Company "))
    assert companies == [f"Company {c}" for c in "ABCD"[:len(companies)]]   # no gaps
    assert consistency_check(res)["all_consistent"]


def test_role_labels_off_gives_letters():
    docs = [(p.name, p.read_text()) for p in sorted((ROOT / "samples").glob("*.txt"))]
    _, res = run(docs, role_labels=False)
    assert all(e.label.split()[0] in ("Person", "Company", "Kept") for e in res.entities
               if e.type in ("PERSON", "ORGANIZATION"))


def test_rename_label_everywhere_and_unique():
    docs = [(p.name, p.read_text()) for p in sorted((ROOT / "samples").glob("*.txt"))]
    s, res = run(docs, role_labels=False)
    northgate = entity(res, "Northgate")
    s.rename(northgate.key, "Receiving broker")
    res2 = s.run()
    assert all("Receiving broker" in out for i, out in enumerate(res2.outputs)
               if any(m[0] == i for m in entity(res2, "Northgate").mentions))
    other = entity(res2, "Harborview")
    with pytest.raises(ValueError):
        s.rename(other.key, "receiving broker")      # labels stay unique (case-insensitive)


# ---------------------------------------------------------------- tickers, share prices, links
def test_bare_ticker_from_sec_list_is_redacted_but_not_ordinary_words():
    txt = "Common stock MSFT NASDAQ. Shares of MSFT fell. IT and ALL staff ON the floor were told NOW."
    _, res = run([("d.txt", txt)])
    out = res.outputs[0]
    assert "MSFT" not in out
    for word in ("IT", "ALL", "ON", "NOW"):
        assert re.search(rf"\b{word}\b", out), word


def test_ticker_links_to_its_company():
    _, res = run([("d.txt", "Analysts rate MSFT a buy. Microsoft shipped a new product.")], public_companies="redact")
    assert "Microsoft" not in res.outputs[0]


def test_share_prices_are_redacted_with_context():
    txt = ("Common stock, $0.00000625 par value per share. The offer was $95.00 per share. "
           "The stock fell 9% on the news. It closed at $296.03. Revenue was $51.7 billion.")
    _, res = run([("d.txt", txt)])
    out = res.outputs[0]
    for price in ("$0.00000625", "$95.00", "$296.03"):
        assert price not in out, price
    assert "~$50 billion" in out              # a headline figure is rounded, not treated as a price


def test_record_links_redacted_general_links_reviewed():
    rec = "https://www.sec.gov/Archives/edgar/data/789019/000119312522010667/d301984d8k.htm"
    txt = f"See {rec} for the filing, and https://www.sec.gov/investor for general guidance."
    _, res = run([("d.txt", txt)])
    links = {e.canonical: e for e in res.entities if e.type == "URL"}
    assert links[rec].status == "auto" and rec not in res.outputs[0]
    general = links["https://www.sec.gov/investor"]
    assert general.status == "pending"
    assert "guidance." in res.outputs[0]       # the sentence's full stop isn't swallowed by the link


def test_address_uses_document_caption_and_zip_is_redacted():
    data = (FIX / "msft_8k_2022-01-18.htm").read_bytes()
    s = Session.from_files([("8k.htm", data)])
    res = s.run()
    out = res.outputs[0]
    assert "Redmond" not in out and "98052" not in out
    assert "Address of principal executive offices" in out


def test_renaming_an_address_overrides_city_state():
    data = (FIX / "msft_8k_2022-01-18.htm").read_bytes()
    s = Session.from_files([("8k.htm", data)])
    s.settings["role_labels"] = False
    addr = next(e for e in s.run().entities if e.type == "STREET_ADDRESS")
    assert addr.proposed.startswith("Redmond")          # city/state by default when context labels are off
    s.rename(addr.key, "Address 1")
    out = s.run().outputs[0]
    assert "Address 1" in out and "Redmond" not in out


# ---------------------------------------------------------------- Activision 8-K excerpt
@pytest.fixture(scope="module")
def atvi():
    return run([("atvi.txt", (FIX / "atvi_8k_excerpt.txt").read_text())])


def test_full_company_suffix_is_covered(atvi):
    _, res = atvi
    out = res.outputs[0]
    assert "oration" not in out and "Microsoft" not in out
    assert "INC." not in out and "Activision" not in out.title()


def test_address_with_building_is_one_redaction(atvi):
    _, res = atvi
    out = res.outputs[0]
    for part in ("2701", "Olympic", "Building B", "Santa Monica", "90404"):
        assert part not in out, part


def test_delisted_ticker_found_by_context(atvi):
    _, res = atvi
    assert "ATVI" not in res.outputs[0]


def test_signatory_title_is_redacted(atvi):
    _, res = atvi
    out = res.outputs[0]
    assert "Corporate Governance" not in out and "Corporate Secretary" not in out
    assert "Ellenberg" not in out


@pytest.mark.parametrize("style,marker", [("redacted", "[REDACTED]"), ("remove", None)])
def test_placeholder_styles(style, marker):
    s, res = run([("atvi.txt", (FIX / "atvi_8k_excerpt.txt").read_text())], placeholder_style=style)
    out = res.outputs[0]
    assert "Microsoft" not in out and "Company A" not in out and "Person A" not in out
    if marker:
        assert marker in out
    else:
        assert '("")' not in out and "“”" not in out
    assert "October 2023" in out               # generalized values stay readable
    assert consistency_check(res)["leak_free"]
